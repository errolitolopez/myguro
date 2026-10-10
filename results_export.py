"""Export the results of an assignment (a sheet given to a class) to Excel or PDF.

Both use the same data as the results page (app.results_data). Excel has every column
as numbers so teachers can compute grades; the PDF is an A4 landscape page to print or send.
"""

import io
import os
from datetime import datetime, timezone

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FONTS_DIR = os.path.join(os.path.dirname(__file__), "fonts")

STATUS = {"not_started": "Not started", "in_progress": "Answering", "submitted": "Submitted"}


def local_time(utc_text, timezone_info):
    """"2026-10-10 08:35:08" (UTC, from SQLite) -> "2026-10-10 16:35" in the school's time zone."""
    if not utc_text:
        return ""
    moment = datetime.strptime(utc_text, "%Y-%m-%d %H:%M:%S").replace(tzinfo=timezone.utc)
    return moment.astimezone(timezone_info).strftime("%Y-%m-%d %H:%M")


def percent(score, total_points):
    """Earned points as a percent of the sheet's total points ("" when not submitted)."""
    if not score or not total_points:
        return ""
    return round(100 * score["earned"] / total_points)


def status_text(student):
    """"Submitted", "Submitted (paper, late)", "Not started", ..."""
    extras = [word for word, flag in (("paper", student["paper"]), ("late", student["late"])) if flag]
    return STATUS[student["status"]] + (f" ({', '.join(extras)})" if extras else "")


def details_lines(data, teacher, exported_at):
    assignment = data["assignment"]
    summary = data["summary"]
    lines = [
        " · ".join(filter(None, [assignment["class_name"], assignment["subject"], f"Teacher: {teacher}"])),
        f"Due: {assignment['due_date'] or 'no due date'} · Total points: {assignment['total_points']} · "
        f"Exported: {exported_at}",
        f"Submitted: {summary['submitted']} of {summary['students']} · Not started: {summary['not_started']} · "
        f"Late: {summary['late']} · Points still to check: {summary['to_check']}",
    ]
    return lines


# --- Excel --------------------------------------------------------------------------

def results_xlsx_bytes(data, teacher, exported_at, timezone_info):
    """The results as an .xlsx file (bytes): a title block, then one row per student."""
    assignment = data["assignment"]
    sections = data["sections"]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Results"

    sheet.append([assignment["title"]])
    sheet["A1"].font = Font(bold=True, size=14)
    for line in details_lines(data, teacher, exported_at):
        sheet.append([line])
    sheet.append([])

    headers = ["No.", "Student", "Username", "Grade", "Section", "Status", "Submitted", "Tries",
               "Score", "Total points", "To check", "% of total"]
    headers += [f"{section['title']} ({section['points']})" for section in sections]
    sheet.append(headers)
    header_row = sheet.max_row
    for cell in sheet[header_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E9ECEF")
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for number, student in enumerate(data["students"], start=1):
        score = student["score"]
        row = [
            number, student["name"], student["username"], student["grade"], student["section"],
            status_text(student), local_time(student["submitted_at"], timezone_info), student["attempts"] or "",
            score["earned"] if score else "", assignment["total_points"] if score else "",
            score["pending"] if score else "", percent(score, assignment["total_points"]),
        ]
        for section in sections:
            part = student["section_scores"].get(str(section["id"]))
            row.append(part["earned"] if part else "")
        sheet.append(row)

    widths = [5, 28, 16, 10, 12, 24, 17, 6, 7, 8, 8, 9] + [14] * len(sections)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=3)  # keep headers and names visible

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


# --- PDF ----------------------------------------------------------------------------

BLACK = (33, 37, 41)
GREY = (108, 117, 125)
LINE = 6  # mm, height of a table row


class ResultsPDF(FPDF):
    def __init__(self):
        super().__init__(orientation="landscape", unit="mm", format="A4")
        self.set_margins(10, 10, 10)
        self.set_auto_page_break(auto=True, margin=12)
        self.add_font("DejaVu", "", os.path.join(FONTS_DIR, "DejaVuSans.ttf"))
        self.add_font("DejaVu", "B", os.path.join(FONTS_DIR, "DejaVuSans-Bold.ttf"))
        self.header_cells = []  # repeated at the top of each new page

    def text_style(self, size=8, bold=False, grey=False):
        self.set_font("DejaVu", "B" if bold else "", size)
        self.set_text_color(*(GREY if grey else BLACK))

    def fit(self, text, width):
        """Shorten text with "…" so it fits in a cell of this width."""
        text = str(text)
        if self.get_string_width(text) <= width - 2:
            return text
        while text and self.get_string_width(text + "…") > width - 2:
            text = text[:-1]
        return text + "…"

    def table_row(self, cells, bold=False, fill=False):
        """cells: [(text, width, align)]. Student rows are one line (long text is shortened);
        the header row (bold) wraps onto up to 3 lines so section titles stay readable."""
        self.text_style(7.5, bold)
        lines = 1
        if bold:
            lines = min(3, max(len(self.multi_cell(width, 4, str(text), dry_run=True, output="LINES"))
                               for text, width, _ in cells))
        height = LINE if lines == 1 else 4 * lines + 2
        if self.will_page_break(height):
            self.add_page()
            if self.header_cells and not bold:
                self.table_row(self.header_cells, bold=True, fill=True)
            self.text_style(7.5, bold)
        self.set_fill_color(233, 236, 239)
        self.set_draw_color(206, 212, 218)
        x, y = self.get_x(), self.get_y()
        for text, width, align in cells:
            if lines == 1:
                self.cell(width, height, self.fit(text, width), border=1, align=align, fill=fill)
            else:
                self.rect(x, y, width, height, style="DF" if fill else "D")
                self.set_xy(x, y + 1)
                self.multi_cell(width, 4, str(text), align=align, max_line_height=4)
                x += width
                self.set_xy(x, y)
        self.set_xy(self.l_margin, y + height)

    def footer(self):
        self.set_y(-9)
        self.text_style(7, grey=True)
        self.cell(0, 4, f"Page {self.page_no()}", align="R")


def results_pdf_bytes(data, teacher, exported_at, timezone_info):
    """The results as a PDF (bytes): title, details, then a table of students."""
    assignment = data["assignment"]
    sections = data["sections"]
    pdf = ResultsPDF()
    pdf.set_title(f"{assignment['title']} - {assignment['class_name']}")
    pdf.add_page()

    pdf.text_style(14, bold=True)
    pdf.multi_cell(0, 7, assignment["title"], new_x="LMARGIN", new_y="NEXT")
    pdf.text_style(8.5, grey=True)
    for line in details_lines(data, teacher, exported_at):
        pdf.cell(0, 4.5, line, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)

    # Fixed columns, then the section columns share what's left of the page width.
    fixed = [("No.", 9, "R"), ("Student", 52, "L"), ("Grade & section", 32, "L"), ("Status", 40, "L"),
             ("Submitted", 27, "L"), ("Score", 18, "R"), ("%", 11, "R")]
    left = pdf.epw - sum(width for _, width, _ in fixed)
    section_width = left / len(sections) if sections else 0
    if not sections:  # no section columns: give the space to the student's name
        fixed[1] = ("Student", 52 + left, "L")
    header = fixed + [(f"{section['title']} ({section['points']})", section_width, "C") for section in sections]
    pdf.header_cells = header
    pdf.table_row([(title, width, "L" if align == "L" else "C") for title, width, align in header],
                  bold=True, fill=True)

    for number, student in enumerate(data["students"], start=1):
        score = student["score"]
        score_text = ""
        if score:
            score_text = f"{score['earned']}/{assignment['total_points']}"
            if score["pending"]:
                score_text += "*"
        place = " - ".join(filter(None, [student["grade"], student["section"]]))
        row = [
            (number, fixed[0][1], "R"), (student["name"], fixed[1][1], "L"), (place, fixed[2][1], "L"),
            (status_text(student), fixed[3][1], "L"), (local_time(student["submitted_at"], timezone_info), fixed[4][1], "L"),
            (score_text, fixed[5][1], "R"), (percent(score, assignment["total_points"]), fixed[6][1], "R"),
        ]
        for section in sections:
            part = student["section_scores"].get(str(section["id"]))
            text = "" if not part else f"{part['earned']}" + ("*" if part["pending"] else "")
            row.append((text, section_width, "C"))
        pdf.table_row(row)

    if any(student["score"] and student["score"]["pending"] for student in data["students"]):
        pdf.ln(2)
        pdf.text_style(7.5, grey=True)
        pdf.cell(0, 4, "* Has short answers the teacher hasn't checked yet.", new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())
