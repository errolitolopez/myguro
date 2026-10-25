"""Reading an uploaded Word (.docx) activity sheet into lines, and the template to fill in.

Only the reading is here (lines of text with their heading level); the rules for a sheet are in
domain/docx_import.py.
"""

import io
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.exceptions import PackageNotFoundError
from docx.shared import Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph

from myguro.config import GENERATED_BY
from myguro.errors import ValidationError

NAME_TABLE_START = "name of student"  # the Name / Grade / Date box at the top of a LAS is skipped


def read_docx_lines(content):
    """[(kind, text)] from a .docx file (bytes) in reading order. kind is "h1" (Heading 1), "h2"
    (Heading 2) or "p". Empty lines are skipped; a table row becomes one line, cells joined by " | "."""
    try:
        document = Document(io.BytesIO(content))
    except (PackageNotFoundError, zipfile.BadZipFile, KeyError, ValueError, OSError):
        raise ValidationError("This is not a Word .docx file. In Word, use Save As → "
                              "Word Document (.docx).") from None
    lines = []
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            paragraph = Paragraph(child, document)
            text = " ".join(paragraph.text.split())
            style = paragraph.style.name if paragraph.style is not None else ""
            if text:
                lines.append(({"Heading 1": "h1", "Heading 2": "h2"}.get(style, "p"), text))
        elif child.tag.endswith("}tbl"):
            table = Table(child, document)
            rows = [[" ".join(cell.text.split()) for cell in row.cells] for row in table.rows]
            if any(cell.casefold().startswith(NAME_TABLE_START) for row in rows for cell in row):
                continue
            for row in rows:
                cells = []
                for cell in row:
                    if cell and cell not in cells:  # merged cells repeat their text
                        cells.append(cell)
                if cells:
                    lines.append(("p", " | ".join(cells)))
    return lines


EXAMPLE = [
    ("h1", "I. Lesson"),
    ("p", "Text sections have no Type line. Write the lesson here, one paragraph per line."),
    ("h1", "II. Pre-test"),
    ("p", "Type: Multiple Choice"),
    ("p", "Points: 1"),
    ("p", "Directions: Choose the letter of the correct answer."),
    ("p", "1. Which tradition originated in India?"),
    ("p", "A. Confucianism"),
    ("p", "B. Hinduism"),
    ("p", "C. Taoism"),
    ("p", "2. What is the central concern of Buddhism?"),
    ("p", "A. Social hierarchy"),
    ("p", "B. Suffering and liberation"),
    ("p", "C. Military strategy"),
    ("h1", "III. True or False"),
    ("p", "Type: True or False"),
    ("p", "Directions: Write True if the statement is correct, False if not."),
    ("p", "1. Wu wei means doing nothing at all."),
    ("p", "2. Karma is about action and its moral consequences."),
    ("h1", "IV. Identification"),
    ("p", "Type: Identification"),
    ("p", "Directions: Write the answer."),
    ("p", "1. Which tradition teaches the Four Noble Truths?"),
    ("p", "2. Which tradition teaches living in harmony with the Tao?"),
    ("h1", "V. Fill in the Blank"),
    ("p", "Type: Fill in the Blank"),
    ("p", "Directions: Write the missing words. Each ___ is one blank."),
    ("p", "1. Hinduism started in ___ and Confucianism started in ___."),
    ("h1", "VI. Matching"),
    ("p", "Type: Matching"),
    ("p", "Points: 4"),
    ("p", "Directions: Match each concept in Column A to its meaning in Column B."),
    ("p", "1. Ren"),
    ("p", "2. Dharma"),
    ("p", "3. Wu wei"),
    ("p", "A. Effortless action"),
    ("p", "B. Ethical duty"),
    ("p", "C. Humaneness"),
    ("p", "D. Cycle of rebirth"),
    ("h1", "VII. Reflection"),
    ("p", "Type: Short Answer"),
    ("p", "Points: 5"),
    ("p", "Directions: Answer in 2 to 3 sentences."),
    ("p", "1. Which tradition speaks most to you, and why?"),
    ("h1", "ANSWER KEY"),
    ("h2", "II. Pre-test"),
    ("p", "1. B (Hinduism)"),
    ("p", "2. B"),
    ("h2", "III. True or False"),
    ("p", "1. FALSE"),
    ("p", "2. TRUE"),
    ("h2", "IV. Identification"),
    ("p", "1. Buddhism"),
    ("p", "2. Taoism / Daoism"),
    ("h2", "V. Fill in the Blank"),
    ("p", "1. India; China"),
    ("h2", "VI. Matching"),
    ("p", "1. C"),
    ("p", "2. B"),
    ("p", "3. A"),
    ("h2", "VII. Reflection"),
    ("p", "1. Any tradition, if the reasons come from the lesson. (optional guide for checking)"),
]

INSTRUCTIONS = [
    "How to use this template",
    "1. Keep each section heading as Heading 1 (Home > Styles > Heading 1). Everything under a heading "
    "belongs to that section.",
    "2. A section with no Type line is a text section (lesson, objectives, discussion).",
    "3. A section whose first line is Type: ... is an activity. Types: Multiple Choice, True or False, "
    "Identification, Fill in the Blank, Matching, Short Answer.",
    "4. Optional lines: Points: 2 (for every question of the section; default 1, Short Answer 5) and "
    "Directions: ... (shown above the questions).",
    "5. Number every question: 1. 2. 3. Multiple choice choices are A. B. C. lines under the question. "
    "Fill in the blank uses ___ for each blank. A Matching section is one question: numbered lines are "
    "Column A, lettered lines (A. B. C.) are Column B.",
    "6. The last section is Heading 1 ANSWER KEY. Repeat each activity's heading (Heading 2) exactly as it is "
    "written above, including numbers and anything in brackets, then one line "
    "per question: 1. B  |  2. TRUE  |  3. Taoism / Daoism  |  4. India; China. Matching: the Column B "
    "letter for each Column A number. Short Answer: an optional guide for checking. Put spaces around the "
    "slash between accepted answers (Taoism / Daoism); a slash without spaces, like 1/2, is part of the answer.",
    "7. Anything before the first Heading 1 (like this page) is not imported. Delete the example below "
    "and write your own.",
]


def add_footer_note(document, text):
    """A small gray line centered in the page footer of every page. The importer reads only the
    body, so the note never becomes part of an imported activity."""
    for section in document.sections:
        paragraph = section.footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(text)
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x6C, 0x75, 0x7D)


def sheet_template_docx_bytes():
    """The Word file to fill in: how to use it, then an example of every type and its answer key."""
    document = Document()
    add_footer_note(document, GENERATED_BY)
    document.add_heading("MyGURO activity template", level=0)
    document.add_paragraph().add_run(INSTRUCTIONS[0]).bold = True
    for text in INSTRUCTIONS[1:]:
        document.add_paragraph(text)
    for kind, text in EXAMPLE:
        if kind == "h1":
            document.add_heading(text, level=1)
        elif kind == "h2":
            document.add_heading(text, level=2)
        else:
            document.add_paragraph(text)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


TYPE_NAMES = {"multiple_choice": "Multiple Choice", "true_false": "True or False", "identification": "Identification",
              "fill_blank": "Fill in the Blank", "matching": "Matching", "short_answer": "Short Answer"}
LETTERS = "ABCDEFGHIJKL"


def question_blocks(questions):
    """The questions split into runs of one type: a Word section has one Type, and a matching
    question is a section of its own. [[question, ...], ...]"""
    blocks = []
    for question in questions:
        if blocks and question["type"] == blocks[-1][0]["type"] != "matching":
            blocks[-1].append(question)
        else:
            blocks.append([question])
    return blocks


def key_line(question):
    """The answer-key text of a question, as docx_import reads it."""
    answer, choices = question["answer"], question["choices"]
    kind = question["type"]
    if kind == "multiple_choice":
        return f"{LETTERS[answer]} ({choices[answer]})"
    if kind == "true_false":
        return "TRUE" if answer else "FALSE"
    if kind == "identification":
        return " / ".join(answer)
    if kind == "fill_blank":
        return "; ".join(" / ".join(blank) for blank in answer)
    return answer or ""  # short answer: the guide for checking


def sheet_docx_lines(activity):
    """[(kind, text)] of an activity in the import template format (see EXAMPLE): sections, then
    the ANSWER KEY. A section that mixes question types becomes one Word section per run of the
    same type ("Title", "Title (part 2)"), and every heading is made unique, because the answer
    key finds a section by its heading."""
    from myguro.domain.richtext import rich_to_text

    lines, key, used = [], [], set()

    def heading(title):
        name, count = title, 1
        while name.casefold() in used:
            count += 1
            name = f"{title} ({count})"
        used.add(name.casefold())
        return name

    for number, section in enumerate(activity["sections"], start=1):
        description = rich_to_text(section["description"]) if section["description_format"] == "rich" \
            else section["description"]
        title = section["title"] or f"Section {number}"
        if section["kind"] == "text":
            lines.append(("h1", heading(title)))
            lines.extend(("p", line) for line in description.splitlines() if line.strip())
            continue
        for part, block in enumerate(question_blocks(section["questions"]), start=1):
            name = heading(title if part == 1 else f"{title} (part {part})")
            kind = block[0]["type"]
            points = max((question["points"] for question in block), key=[q["points"] for q in block].count)
            lines += [("h1", name), ("p", f"Type: {TYPE_NAMES[kind]}"), ("p", f"Points: {points}")]
            key.append(("h2", name))
            directions = block[0]["prompt"] if kind == "matching" else description if part == 1 else ""
            for index, line in enumerate(directions.splitlines()):
                if line.strip():
                    lines.append(("p", f"Directions: {line}" if index == 0 else line))
            if kind == "matching":
                question = block[0]
                lines += [("p", f"{item}. {pair['item']}") for item, pair in enumerate(question["answer"], start=1)]
                lines += [("p", f"{LETTERS[index]}. {choice}") for index, choice in enumerate(question["choices"])]
                key += [("p", f"{item}. {LETTERS[pair['match']]}") for item, pair in enumerate(question["answer"], start=1)]
                continue
            for item, question in enumerate(block, start=1):
                lines.append(("p", f"{item}. {question['prompt']}"))
                lines += [("p", f"{LETTERS[index]}. {choice}") for index, choice in enumerate(question["choices"] or [])]
                answer = key_line(question)
                if answer:
                    key.append(("p", f"{item}. {answer}"))
    return lines + [("h1", "ANSWER KEY")] + key


def sheet_docx_bytes(activity, teacher):
    """An activity as a Word file in the import template format, with its ANSWER KEY: the teacher's
    copy to keep or edit, which can also be imported as a new activity."""
    document = Document()
    add_footer_note(document, GENERATED_BY)
    document.add_heading(activity["title"], level=0)
    document.add_paragraph("  ·  ".join(filter(None, [activity["subject"], teacher])))
    document.add_paragraph("Teacher's copy with the answer key. To use it again in MyGURO: Activities → "
                           "Import from Word. Keep the headings and the Type, Points and "
                           "Directions lines as they are. (This part is not imported.)")
    for kind, text in sheet_docx_lines(activity):
        if kind == "h1":
            document.add_heading(text, level=1)
        elif kind == "h2":
            document.add_heading(text, level=2)
        else:
            document.add_paragraph(text)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()
