"""Export the results of an assignment (a sheet given to a class) to Excel or PDF.

Both use the same data as the results page (app.results_data). Excel has every column
as numbers so teachers can compute grades; the PDF is an A4 landscape page to print or send.
"""

import io
import os
from datetime import datetime, timezone

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from myguro.config import GENERATED_BY
from myguro.domain.richtext import rich_to_text

FONTS_DIR = os.path.join(os.path.dirname(__file__), "fonts")

STATUS = {"not_started": "Not started", "in_progress": "Answering", "submitted": "Submitted",
          "excluded": "Excluded"}


def local_time(utc_text, timezone_info):
    """"2026-10-10 08:35:08" (UTC, from SQLite) -> "2026-10-10 16:35" in the school's time zone."""
    if not utc_text:
        return ""
    moment = datetime.strptime(utc_text, "%Y-%m-%d %H:%M:%S").replace(tzinfo=timezone.utc)
    return moment.astimezone(timezone_info).strftime("%Y-%m-%d %H:%M")


def percent(score, total_points):
    """Earned points as a percent of the sheet's total points ("" when not submitted)."""
    if not score or not total_points:
        return ""
    return round(100 * score["earned"] / total_points)


def status_text(student):
    """"Submitted", "Submitted (paper, late)", "Not started", ..."""
    extras = [word for word, flag in (("paper", student["paper"]), ("late", student["late"])) if flag]
    return STATUS[student["status"]] + (f" ({', '.join(extras)})" if extras else "")


def details_lines(data, teacher, exported_at):
    assignment = data["assignment"]
    summary = data["summary"]
    lines = [
        " · ".join(filter(None, [assignment["class_name"], assignment["subject"], f"Teacher: {teacher}"])),
        f"Due: {assignment['due_date'] or 'no due date'} · Total points: {assignment['total_points']} · "
        f"Exported: {exported_at}",
        f"Submitted: {summary['submitted']} of {summary['students']} · Not started: {summary['not_started']} · "
        f"Late: {summary['late']} · Points still to check: {summary['to_check']}",
    ]
    return lines


# --- Excel --------------------------------------------------------------------------

def add_generated_by(sheet):
    """The note under the table (visible) and in the printed footer, after everything else is written."""
    sheet.append([])
    sheet.append([GENERATED_BY])
    sheet.cell(row=sheet.max_row, column=1).font = Font(italic=True, color="6C757D")
    sheet.oddFooter.center.text = GENERATED_BY


LOGIN_DETAILS_COLUMNS = ["Name", "Username", "Password", "Grade", "Section"]


def login_details_xlsx_bytes(students):
    """The students' login details as an .xlsx file (bytes), to give out or print. Each student is a
    dict with name, username, password, grade and section. Everything is text, so a password like
    000123 keeps its zeros."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Login details"
    sheet.append(LOGIN_DETAILS_COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E9ECEF")
    for student in students:
        sheet.append([student["name"], student["username"], student["password"], student["grade"],
                      student["section"]])
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.number_format = "@"
            cell.data_type = "s"
    for index, width in enumerate((30, 18, 16, 12, 14), start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    add_generated_by(sheet)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def results_xlsx_bytes(data, teacher, exported_at, timezone_info):
    """The results as an .xlsx file (bytes): a title block, then one row per student."""
    assignment = data["assignment"]
    sections = data["sections"]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Results"

    sheet.append([assignment["title"]])
    sheet["A1"].font = Font(bold=True, size=14)
    for line in details_lines(data, teacher, exported_at):
        sheet.append([line])
    sheet.append([])

    headers = ["No.", "Student", "Username", "Grade", "Section", "Status", "Submitted", "Tries",
               "Score", "Total points", "To check", "% of total"]
    headers += [f"{section['title']} ({section['points']})" for section in sections]
    sheet.append(headers)
    header_row = sheet.max_row
    for cell in sheet[header_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E9ECEF")
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for number, student in enumerate(data["students"], start=1):
        score = student["score"]
        row = [
            number, student["name"], student["username"], student["grade"], student["section"],
            status_text(student), local_time(student["submitted_at"], timezone_info), student["attempts"] or "",
            score["earned"] if score else "", assignment["total_points"] if score else "",
            score["pending"] if score else "", percent(score, assignment["total_points"]),
        ]
        for section in sections:
            part = student["section_scores"].get(str(section["id"]))
            row.append(part["earned"] if part else "")
        sheet.append(row)

    widths = [5, 28, 16, 10, 12, 24, 17, 6, 7, 8, 8, 9] + [14] * len(sections)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=3)  # keep headers and names visible
    add_generated_by(sheet)

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


# --- PDF ----------------------------------------------------------------------------

BLACK = (33, 37, 41)
GREY = (108, 117, 125)
LINE = 6  # mm, height of a table row


class ResultsPDF(FPDF):
    def __init__(self):
        super().__init__(orientation="landscape", unit="mm", format="A4")
        self.set_margins(10, 10, 10)
        self.set_auto_page_break(auto=True, margin=12)
        self.add_font("DejaVu", "", os.path.join(FONTS_DIR, "DejaVuSans.ttf"))
        self.add_font("DejaVu", "B", os.path.join(FONTS_DIR, "DejaVuSans-Bold.ttf"))
        self.header_cells = []  # repeated at the top of each new page

    def text_style(self, size=8, bold=False, grey=False):
        self.set_font("DejaVu", "B" if bold else "", size)
        self.set_text_color(*(GREY if grey else BLACK))

    def fit(self, text, width):
        """Shorten text with "…" so it fits in a cell of this width."""
        text = str(text)
        if self.get_string_width(text) <= width - 2:
            return text
        while text and self.get_string_width(text + "…") > width - 2:
            text = text[:-1]
        return text + "…"

    def table_row(self, cells, bold=False, fill=False):
        """cells: [(text, width, align)]. Student rows are one line (long text is shortened);
        the header row (bold) wraps onto up to 3 lines so section titles stay readable."""
        self.text_style(7.5, bold)
        lines = 1
        if bold:
            lines = min(3, max(len(self.multi_cell(width, 4, str(text), dry_run=True, output="LINES"))
                               for text, width, _ in cells))
        height = LINE if lines == 1 else 4 * lines + 2
        if self.will_page_break(height):
            self.add_page()
            if self.header_cells and not bold:
                self.table_row(self.header_cells, bold=True, fill=True)
            self.text_style(7.5, bold)
        self.set_fill_color(233, 236, 239)
        self.set_draw_color(206, 212, 218)
        x, y = self.get_x(), self.get_y()
        for text, width, align in cells:
            if lines == 1:
                self.cell(width, height, self.fit(text, width), border=1, align=align, fill=fill)
            else:
                self.rect(x, y, width, height, style="DF" if fill else "D")
                self.set_xy(x, y + 1)
                self.multi_cell(width, 4, str(text), align=align, max_line_height=4)
                x += width
                self.set_xy(x, y)
        self.set_xy(self.l_margin, y + height)

    def footer(self):
        self.set_y(-9)
        self.text_style(7, grey=True)
        self.cell(0, 4, GENERATED_BY, align="C")
        self.set_x(self.l_margin)  # the same line again, for the page number on the right
        self.cell(0, 4, f"Page {self.page_no()}", align="R")


def results_pdf_bytes(data, teacher, exported_at, timezone_info):
    """The results as a PDF (bytes): title, details, then a table of students."""
    assignment = data["assignment"]
    sections = data["sections"]
    pdf = ResultsPDF()
    pdf.set_title(f"{assignment['title']} - {assignment['class_name']}")
    pdf.add_page()

    pdf.text_style(14, bold=True)
    pdf.multi_cell(0, 7, assignment["title"], new_x="LMARGIN", new_y="NEXT")
    pdf.text_style(8.5, grey=True)
    for line in details_lines(data, teacher, exported_at):
        pdf.cell(0, 4.5, line, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)

    # Fixed columns, then the section columns share what's left of the page width.
    fixed = [("No.", 9, "R"), ("Student", 52, "L"), ("Grade & section", 32, "L"), ("Status", 40, "L"),
             ("Submitted", 27, "L"), ("Score", 18, "R"), ("%", 11, "R")]
    left = pdf.epw - sum(width for _, width, _ in fixed)
    section_width = left / len(sections) if sections else 0
    if not sections:  # no section columns: give the space to the student's name
        fixed[1] = ("Student", 52 + left, "L")
    header = fixed + [(f"{section['title']} ({section['points']})", section_width, "C") for section in sections]
    pdf.header_cells = header
    pdf.table_row([(title, width, "L" if align == "L" else "C") for title, width, align in header],
                  bold=True, fill=True)

    for number, student in enumerate(data["students"], start=1):
        score = student["score"]
        score_text = ""
        if score:
            score_text = f"{score['earned']}/{assignment['total_points']}"
            if score["pending"]:
                score_text += "*"
        place = " - ".join(filter(None, [student["grade"], student["section"]]))
        row = [
            (number, fixed[0][1], "R"), (student["name"], fixed[1][1], "L"), (place, fixed[2][1], "L"),
            (status_text(student), fixed[3][1], "L"), (local_time(student["submitted_at"], timezone_info), fixed[4][1], "L"),
            (score_text, fixed[5][1], "R"), (percent(score, assignment["total_points"]), fixed[6][1], "R"),
        ]
        for section in sections:
            part = student["section_scores"].get(str(section["id"]))
            text = "" if not part else f"{part['earned']}" + ("*" if part["pending"] else "")
            row.append((text, section_width, "C"))
        pdf.table_row(row)

    if any(student["score"] and student["score"]["pending"] for student in data["students"]):
        pdf.ln(2)
        pdf.text_style(7.5, grey=True)
        pdf.cell(0, 4, "* Has short answers the teacher hasn't checked yet.", new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


# --- Class gradebook ----------------------------------------------------------------
# One row per student, one column per activity given to the class (GradebookService.gradebook).

def gradebook_cell_text(cell):
    """What a gradebook cell says: "8/10" ("8/10*" when short answers are still unchecked),
    "Answering", "Excluded", or "" (not started, or not given to this student)."""
    if cell is None or cell["status"] == "not_started":
        return ""
    if cell["status"] == "excluded":
        return "Excluded"
    if cell["earned"] is None:
        return "Answering"
    return f"{cell['earned']}/{cell['possible']}" + ("*" if cell["pending"] else "")


def gradebook_lines(data, teacher, exported_at):
    return [
        " · ".join(filter(None, [data["class"]["name"], data["class"]["subject"], f"Teacher: {teacher}"])),
        f"{len(data['students'])} students · {len(data['activities'])} activities · Exported: {exported_at}",
    ]


def gradebook_xlsx_bytes(data, teacher, exported_at, timezone_info):
    """The gradebook as an .xlsx file (bytes). Scores are numbers (points earned), so teachers can
    compute their own grades; the last columns are how many were submitted and the average %."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Gradebook"
    sheet.append(["Gradebook"])
    sheet["A1"].font = Font(bold=True, size=14)
    for line in gradebook_lines(data, teacher, exported_at):
        sheet.append([line])
    sheet.append([])

    activities = data["activities"]
    headers = ["No.", "Student", "Username", "Grade", "Section"]
    headers += [f"{activity['title']} ({activity['total_points']})" for activity in activities]
    headers += ["Submitted", "Average %"]
    sheet.append(headers)
    header_row = sheet.max_row
    for cell in sheet[header_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E9ECEF")
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for number, student in enumerate(data["students"], start=1):
        row = [number, student["name"], student["username"], student["grade"], student["section"]]
        for cell in student["cells"]:
            if cell is not None and cell["earned"] is not None:
                row.append(cell["earned"])
            else:
                row.append(gradebook_cell_text(cell))
        row += [f"{student['submitted']} of {len([c for c in student['cells'] if c is not None])}",
                student["average"] if student["average"] is not None else ""]
        sheet.append(row)

    sheet.append(["", "Class average %", "", "", ""]
                 + [activity["average"] if activity["average"] is not None else "" for activity in activities])
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)

    widths = [5, 28, 16, 10, 12] + [14] * len(activities) + [11, 11]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=3)  # keep headers and names visible
    add_generated_by(sheet)

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def gradebook_pdf_bytes(data, teacher, exported_at, timezone_info):
    """The gradebook as a PDF (bytes), A4 landscape. The activity columns share the width that
    is left after the student, submitted and average columns."""
    activities = data["activities"]
    pdf = ResultsPDF()
    pdf.set_title(f"Gradebook - {data['class']['name']}")
    pdf.add_page()

    pdf.text_style(14, bold=True)
    pdf.multi_cell(0, 7, "Gradebook", new_x="LMARGIN", new_y="NEXT")
    pdf.text_style(8.5, grey=True)
    for line in gradebook_lines(data, teacher, exported_at):
        pdf.cell(0, 4.5, line, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)

    fixed_start = [("No.", 9, "R"), ("Student", 52, "L"), ("Grade & section", 32, "L")]
    fixed_end = [("Done", 14, "C"), ("Avg %", 13, "R")]
    left = pdf.epw - sum(width for _, width, _ in fixed_start + fixed_end)
    column = max(left / len(activities), 12) if activities else 0
    if not activities:
        fixed_start[1] = ("Student", 52 + left, "L")
    header = (fixed_start
              + [(f"{activity['title']} ({activity['total_points']})", column, "C") for activity in activities]
              + fixed_end)
    pdf.header_cells = header
    pdf.table_row([(title, width, "L" if align == "L" else "C") for title, width, align in header],
                  bold=True, fill=True)

    for number, student in enumerate(data["students"], start=1):
        place = " - ".join(filter(None, [student["grade"], student["section"]]))
        given = len([cell for cell in student["cells"] if cell is not None])
        row = [(number, fixed_start[0][1], "R"), (student["name"], fixed_start[1][1], "L"),
               (place, fixed_start[2][1], "L")]
        row += [(gradebook_cell_text(cell), column, "C") for cell in student["cells"]]
        row += [(f"{student['submitted']}/{given}", fixed_end[0][1], "C"),
                ("" if student["average"] is None else student["average"], fixed_end[1][1], "R")]
        pdf.table_row(row)

    average_row = [("", fixed_start[0][1], "R"), ("Class average %", fixed_start[1][1], "L"),
                   ("", fixed_start[2][1], "L")]
    average_row += [("" if activity["average"] is None else activity["average"], column, "C")
                    for activity in activities]
    average_row += [("", fixed_end[0][1], "C"), ("", fixed_end[1][1], "R")]
    pdf.table_row(average_row, bold=True, fill=True)

    if any(cell and cell["pending"] for student in data["students"] for cell in student["cells"]):
        pdf.ln(2)
        pdf.text_style(7.5, grey=True)
        pdf.cell(0, 4, "* Has short answers the teacher hasn't checked yet.", new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


# --- Attendance ------------------------------------------------------------------------

ATTENDANCE_LETTERS = {"present": "P", "late": "L", "absent": "A", "excused": "E"}
ATTENDANCE_DAYS_PER_BLOCK = 20  # day columns on one page width of the PDF; more days continue below


def attendance_lines(data, teacher, exported_at):
    return [
        " · ".join(filter(None, [data["class"]["name"], data["class"]["subject"], f"Teacher: {teacher}"])),
        f"{len(data['students'])} students · {len(data['days'])} days recorded · Exported: {exported_at}",
    ] + ([f"Days: {range_text(data['range'])}"] if data.get("range") and (data["range"]["from"] or data["range"]["to"]) else [])


def range_text(days):
    """{"from": "2026-10-01", "to": None} -> "from 2026-10-01"; both -> "2026-10-01 to 2026-10-31"."""
    if days["from"] and days["to"]:
        return f"{days['from']} to {days['to']}"
    return f"from {days['from']}" if days["from"] else f"until {days['to']}"


def short_day(isoday):
    """"2026-10-26" -> "10-26"."""
    return isoday[5:]


def attendance_xlsx_bytes(data, teacher, exported_at, timezone_info):
    """The attendance record as an .xlsx file (bytes): a column per recorded day with P (present),
    L (late), A (absent) or E (excused), then the totals as numbers."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Attendance"
    sheet.append(["Attendance"])
    sheet["A1"].font = Font(bold=True, size=14)
    for line in attendance_lines(data, teacher, exported_at):
        sheet.append([line])
    sheet.append(["P = present, L = late, A = absent, E = excused"])
    sheet.append([])

    sheet.append(["No.", "Student", "Username", "Grade", "Section"] + data["days"]
                 + ["Present", "Late", "Absent", "Excused"])
    header_row = sheet.max_row
    for cell in sheet[header_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E9ECEF")
        cell.alignment = Alignment(horizontal="center", vertical="top")

    for number, student in enumerate(data["students"], start=1):
        row = [number, student["name"], student["username"], student["grade"], student["section"]]
        row += [ATTENDANCE_LETTERS.get(student["marks"].get(day), "") for day in data["days"]]
        row += [student["totals"][status] for status in ATTENDANCE_LETTERS]
        sheet.append(row)
        for cell in sheet[sheet.max_row][5:5 + len(data["days"])]:
            cell.alignment = Alignment(horizontal="center")

    widths = [5, 28, 16, 10, 12] + [11] * len(data["days"]) + [9] * 4
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=3)  # keep headers and names visible
    add_generated_by(sheet)

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def attendance_pdf_bytes(data, teacher, exported_at, timezone_info):
    """The attendance record as a PDF (bytes), A4 landscape. Days run across; when there are more
    than ATTENDANCE_DAYS_PER_BLOCK they continue in another table below, with the students again."""
    pdf = ResultsPDF()
    pdf.set_title(f"Attendance - {data['class']['name']}")
    pdf.add_page()

    pdf.text_style(14, bold=True)
    pdf.multi_cell(0, 7, "Attendance", new_x="LMARGIN", new_y="NEXT")
    pdf.text_style(8.5, grey=True)
    for line in attendance_lines(data, teacher, exported_at):
        pdf.cell(0, 4.5, line, new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 4.5, "P = present, L = late, A = absent, E = excused", new_x="LMARGIN", new_y="NEXT")

    start = [("No.", 9, "R"), ("Student", 55, "L")]
    end = [("Pres.", 11, "C"), ("Late", 11, "C"), ("Abs.", 11, "C"), ("Exc.", 11, "C")]
    day_width = 8
    blocks = [data["days"][i:i + ATTENDANCE_DAYS_PER_BLOCK]
              for i in range(0, len(data["days"]), ATTENDANCE_DAYS_PER_BLOCK)] or [[]]
    for days in blocks:
        pdf.ln(3)
        header = start + [(short_day(day), day_width, "C") for day in days] + end
        pdf.header_cells = header
        pdf.table_row(header, bold=True, fill=True)
        for number, student in enumerate(data["students"], start=1):
            row = [(number, start[0][1], "R"), (student["name"], start[1][1], "L")]
            row += [(ATTENDANCE_LETTERS.get(student["marks"].get(day), ""), day_width, "C") for day in days]
            row += [(student["totals"][status], width, "C") for status, (_, width, _) in zip(ATTENDANCE_LETTERS, end)]
            pdf.table_row(row)
    return bytes(pdf.output())


# --- Printable activity sheet ---------------------------------------------------------
# The sheet on paper: A4 portrait, either the student copy (blank answer lines) or the
# teacher's copy with the answer key. Built from ActivityService.get (saved version).

LETTERS = "ABCDEFGHIJKL"
WRITE_LINES = 3  # ruled lines under a short answer


class SheetPDF(FPDF):
    def __init__(self):
        super().__init__(orientation="portrait", unit="mm", format="A4")
        self.set_margins(15, 15, 15)
        self.set_auto_page_break(auto=True, margin=15)
        self.add_font("DejaVu", "", os.path.join(FONTS_DIR, "DejaVuSans.ttf"))
        self.add_font("DejaVu", "B", os.path.join(FONTS_DIR, "DejaVuSans-Bold.ttf"))

    def text_style(self, size=10, bold=False, grey=False):
        self.set_font("DejaVu", "B" if bold else "", size)
        self.set_text_color(*(GREY if grey else BLACK))

    def paragraph(self, text, size=10, bold=False, grey=False, indent=0):
        """A wrapped paragraph; indent moves it right (mm) from the left margin."""
        self.text_style(size, bold, grey)
        self.set_x(self.l_margin + indent)
        self.multi_cell(0, size * 0.5, text, new_x="LMARGIN", new_y="NEXT")

    def footer(self):
        self.set_y(-10)
        self.text_style(8, grey=True)
        self.cell(0, 4, GENERATED_BY, align="C")
        self.set_x(self.l_margin)  # the same line again, for the page number on the right
        self.cell(0, 4, f"Page {self.page_no()}", align="R")


def question_key_text(question):
    """The right answer of a question as one line for the answer key."""
    answer, choices = question["answer"], question["choices"]
    kind = question["type"]
    if kind == "multiple_choice":
        return f"{LETTERS[answer]}. {choices[answer]}"
    if kind == "true_false":
        return "TRUE" if answer else "FALSE"
    if kind == "identification":
        return " / ".join(answer)
    if kind == "fill_blank":
        return "; ".join(" / ".join(blank) for blank in answer)
    if kind == "matching":
        return "; ".join(f"{number}. {LETTERS[pair['match']]}" for number, pair in enumerate(answer, start=1))
    return answer or ""  # short answer: the guide for checking, if any


def print_question(pdf, number, question, with_key):
    """One question: its prompt, the choices or a place to write, and (answer key) the answer."""
    kind = question["type"]
    points = f" ({question['points']} {'pt' if question['points'] == 1 else 'pts'})"
    if pdf.will_page_break(30):
        pdf.add_page()  # keep a question together with its first lines when it is near the page end
    if kind == "matching":
        pdf.paragraph(f"{number}. {question['prompt']}{points}")
        pdf.paragraph("Column A", bold=True, indent=6)
        for item_number, pair in enumerate(question["answer"], start=1):
            pdf.paragraph(f"____ {item_number}. {pair['item']}", indent=6)
        pdf.paragraph("Column B", bold=True, indent=6)
        for letter, choice in zip(LETTERS, question["choices"]):
            pdf.paragraph(f"{letter}. {choice}", indent=6)
    else:
        blank_prefix = "_______ " if kind == "true_false" else ""
        pdf.paragraph(f"{number}. {blank_prefix}{question['prompt']}{points}")
        if kind == "multiple_choice":
            for letter, choice in zip(LETTERS, question["choices"]):
                pdf.paragraph(f"{letter}. {choice}", indent=8)
        elif kind == "identification":
            pdf.paragraph("Answer: ______________________________", indent=8)
        elif kind == "short_answer":
            for _ in range(WRITE_LINES):
                pdf.ln(8)  # room for handwriting
                pdf.line(pdf.l_margin + 8, pdf.get_y(), pdf.w - pdf.r_margin, pdf.get_y())
            pdf.ln(2)
    if with_key:
        answer = question_key_text(question)
        if answer:
            pdf.paragraph(f"Answer: {answer}", size=9, bold=True, indent=8)
    pdf.ln(2)


def sheet_pdf_bytes(activity, teacher, with_key):
    """The activity sheet as a PDF to print (bytes). with_key adds each answer and "ANSWER KEY"."""
    pdf = SheetPDF()
    pdf.set_title(activity["title"])
    pdf.add_page()

    pdf.text_style(15, bold=True)
    pdf.multi_cell(0, 8, activity["title"], new_x="LMARGIN", new_y="NEXT")
    total = sum(question["points"] for section in activity["sections"] for question in section["questions"])
    pdf.paragraph("  ·  ".join(filter(None, [activity["subject"], teacher, f"{total} points"])), size=9, grey=True)
    if with_key:
        pdf.paragraph("ANSWER KEY (teacher's copy)", size=10, bold=True)
    else:
        pdf.ln(2)
        pdf.paragraph(f"Name: {'_' * 36}   Date: {'_' * 16}   Score: ______ / {total}")
    pdf.ln(3)

    for section in activity["sections"]:
        if section["title"]:
            pdf.paragraph(section["title"], size=12, bold=True)
        if section["description"]:
            text = rich_to_text(section["description"]) if section["description_format"] == "rich" \
                else section["description"]
            pdf.paragraph(text, grey=section["kind"] == "activity")
        pdf.ln(2)
        for number, question in enumerate(section["questions"], start=1):
            print_question(pdf, number, question, with_key)
        pdf.ln(2)
    return bytes(pdf.output())
