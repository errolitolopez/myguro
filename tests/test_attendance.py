"""Attendance: a teacher marks each student of a class present, late, absent or excused for a day."""

from datetime import date, timedelta

from conftest import login, make_class

DAY = (date.today() - timedelta(days=1)).isoformat()


def sheet(teacher, class_id, day=DAY):
    return teacher.get(f"/api/classes/{class_id}/attendance?date={day}").get_json()


def test_marks_are_saved_for_the_day_and_counted(teacher):
    class_id, students = make_class(teacher)
    first, second = students[0]["id"], students[1]["id"]
    assert all(row["status"] is None for row in sheet(teacher, class_id)["students"])

    response = teacher.put(f"/api/classes/{class_id}/attendance",
                           json={"day": DAY, "marks": {str(first): "present", str(second): "absent"}})
    assert response.status_code == 200, response.get_json()
    by_id = {row["id"]: row for row in sheet(teacher, class_id)["students"]}
    assert by_id[first]["status"] == "present" and by_id[second]["status"] == "absent"
    assert by_id[second]["totals"] == {"present": 0, "late": 0, "absent": 1, "excused": 0}
    assert sheet(teacher, class_id)["recorded_days"] == [DAY]

    # Saving again replaces the mark; null clears it; another day is separate
    teacher.put(f"/api/classes/{class_id}/attendance",
                json={"day": DAY, "marks": {str(first): "late", str(second): None}})
    by_id = {row["id"]: row for row in sheet(teacher, class_id)["students"]}
    assert by_id[first]["status"] == "late" and by_id[second]["status"] is None
    other_day = (date.today() - timedelta(days=2)).isoformat()
    assert all(row["status"] is None for row in sheet(teacher, class_id, other_day)["students"])


def test_bad_requests_are_refused(teacher):
    class_id, students = make_class(teacher)
    student_id = str(students[0]["id"])
    put = lambda body: teacher.put(f"/api/classes/{class_id}/attendance", json=body)
    assert put({"day": DAY, "marks": {student_id: "sleeping"}}).status_code == 400
    assert put({"day": "not a date", "marks": {}}).status_code == 400
    tomorrow = (date.today() + timedelta(days=2)).isoformat()
    assert put({"day": tomorrow, "marks": {student_id: "present"}}).status_code == 400
    assert put({"day": DAY, "marks": {"999999": "present"}}).status_code == 400


def test_only_the_classs_teacher_can_use_it(teacher, other_teacher):
    class_id, students = make_class(teacher)
    assert other_teacher.get(f"/api/classes/{class_id}/attendance").status_code == 404
    assert other_teacher.put(f"/api/classes/{class_id}/attendance",
                             json={"day": DAY, "marks": {str(students[0]["id"]): "present"}}).status_code == 404
    student = login(students[0]["username"])
    assert student.get(f"/api/classes/{class_id}/attendance").status_code in (401, 403)


def test_the_page_opens_and_the_save_is_logged(teacher):
    class_id, students = make_class(teacher)
    assert teacher.get(f"/classes/{class_id}/attendance").status_code == 200
    teacher.put(f"/api/classes/{class_id}/attendance", json={"day": DAY, "marks": {str(students[0]["id"]): "present"}})
    admin = login("admin", "admin")
    admin.post("/api/change-password", json={"current_password": "admin", "new_password": "adminpass1"})
    entries = admin.get("/api/log?q=attendance").get_json()["entries"]
    assert any(entry["action"] == "attendance.saved" for entry in entries)


def test_the_record_downloads_as_excel_and_pdf(teacher, other_teacher):
    import io

    from openpyxl import load_workbook

    class_id, students = make_class(teacher)
    first = str(students[0]["id"])
    for day, status in ((DAY, "present"), ((date.today() - timedelta(days=2)).isoformat(), "absent")):
        teacher.put(f"/api/classes/{class_id}/attendance", json={"day": day, "marks": {first: status}})

    excel = teacher.get(f"/classes/{class_id}/attendance.xlsx")
    assert excel.status_code == 200 and "Attendance" in excel.headers["Content-Disposition"]
    rows = list(load_workbook(io.BytesIO(excel.data)).active.iter_rows(values_only=True))
    header = next(row for row in rows if row[0] == "No.")
    # A column for every calendar day from the first mark to the last (the day number), marks as letters
    first_day = date.today() - timedelta(days=2)
    assert header[5:7] == (first_day.day, (first_day + timedelta(days=1)).day) and header[7] == "Present"
    student_row = next(row for row in rows if row[1] == students[0]["name"])
    assert student_row[5:7] == ("A", "P") and student_row[-4:] == (1, 0, 1, 0)

    pdf = teacher.get(f"/classes/{class_id}/attendance.pdf")
    assert pdf.status_code == 200 and pdf.data.startswith(b"%PDF")
    # Another teacher's class is not downloadable
    assert other_teacher.get(f"/classes/{class_id}/attendance.xlsx").status_code == 302


def test_the_export_can_be_limited_to_a_range_of_days(teacher):
    import io

    from openpyxl import load_workbook

    class_id, students = make_class(teacher)
    first = str(students[0]["id"])
    days = [(date.today() - timedelta(days=n)).isoformat() for n in (1, 2, 3)]
    for day, status in zip(days, ("present", "absent", "late")):
        teacher.put(f"/api/classes/{class_id}/attendance", json={"day": day, "marks": {first: status}})

    # Both ends are included, and the totals count only the days in the range
    excel = teacher.get(f"/classes/{class_id}/attendance.xlsx?from={days[2]}&to={days[1]}")
    rows = list(load_workbook(io.BytesIO(excel.data)).active.iter_rows(values_only=True))
    header = next(row for row in rows if row[0] == "No.")
    last_two = (date.fromisoformat(days[2]).day, date.fromisoformat(days[1]).day)
    assert header[5:7] == last_two and header[7] == "Present"
    student_row = next(row for row in rows if row[1] == students[0]["name"])
    assert student_row[-4:] == (0, 1, 1, 0)

    assert teacher.get(f"/classes/{class_id}/attendance.pdf?from={days[2]}").status_code == 200
    assert teacher.get(f"/classes/{class_id}/attendance.xlsx?from=oops").status_code == 400
    assert teacher.get(f"/classes/{class_id}/attendance.xlsx?from={days[0]}&to={days[2]}").status_code == 400


def test_the_excel_has_a_column_for_every_day_of_the_range(teacher):
    import io

    from openpyxl import load_workbook

    class_id, students = make_class(teacher)
    teacher.put(f"/api/classes/{class_id}/attendance", json={"day": DAY, "marks": {str(students[0]["id"]): "late"}})
    excel = teacher.get(f"/classes/{class_id}/attendance.xlsx?from=2026-10-01&to=2026-10-31")
    rows = list(load_workbook(io.BytesIO(excel.data)).active.iter_rows(values_only=True))
    header = next(row for row in rows if row[0] == "No.")
    assert header[5:36] == tuple(range(1, 32)) and header[36] == "Present"
    assert next(row for row in rows if row[5] == "Oct 2026")  # the month named above the 1st


def test_saving_one_tap_keeps_the_other_marks_and_logs_the_day_once(teacher):
    class_id, students = make_class(teacher)
    first, second = str(students[0]["id"]), str(students[1]["id"])
    put = lambda marks: teacher.put(f"/api/classes/{class_id}/attendance", json={"day": DAY, "marks": marks})
    # The page saves after every tap: only the student tapped is sent
    put({first: "present"})
    put({second: "absent"})
    put({first: "late"})
    by_id = {row["id"]: row["status"] for row in sheet(teacher, class_id)["students"]}
    assert by_id[students[0]["id"]] == "late" and by_id[students[1]["id"]] == "absent"

    admin = login("admin", "admin")
    admin.post("/api/change-password", json={"current_password": "admin", "new_password": "adminpass1"})
    entries = admin.get("/api/log?q=attendance").get_json()["entries"]
    assert len([entry for entry in entries if entry["action"] == "attendance.saved"]) == 1
