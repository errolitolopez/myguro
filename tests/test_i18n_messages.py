"""Every message the server can send has a Taglish entry (found in the code, so a new one cannot be forgotten)."""

import os

from i18n_messages import messages
from myguro.domain.i18n import translate_message
from myguro.i18n.tl import TAGLISH

# Messages that are not a literal at the raise site: each is covered by its own entry (see tl_messages.py) or is
# built from messages that are. The test lists them so a new one has to be added here on purpose.
KNOWN_DYNAMIC = {
    ("myguro/domain/accounts.py", "NAME_FORMAT_HELP"),  # a constant: "Write the name as Last name, First name ..."
    ("myguro/errors.py", "say(error.message)"), ("myguro/errors.py", "say(api_message)"),  # already translated
    ("myguro/imports.py", "header_message"), ("myguro/imports.py", "empty_message"),
    ("myguro/imports.py", "too_many_message"),  # given by the caller: the student import and its header help
    ("myguro/services/chat.py", "rules.CHAT_FULL_MESSAGE"),  # a constant
    ("myguro/services/results.py", "message"),  # a parameter: "Not found." and the like
    ("myguro/domain/accounts.py", "NAME_FORMAT_HELP"),
    ("myguro/domain/docx_import.py", "message"), ("myguro/domain/docx_import.py", "failure.message"),
    ("myguro/domain/docx_import.py", "failure.message.removeprefix('Question: ')"),  # the messages of clean_question (listed)
    ("myguro/services/schedule.py", "self.describe(refused[0], names)"),  # the schedule's own problem texts
}


def test_every_server_message_has_a_taglish_entry():
    os.chdir(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    found, dynamic = messages()
    # The database upgrade (migrations.py) talks to the person who runs the server, not to a user of the pages.
    missing = sorted(text for text, places in found.items()
                     if text not in TAGLISH and not all(place.startswith("myguro/migrations.py") for place in places))
    assert not missing, f"no Taglish for these server messages: {missing}"
    unknown = [item for item in dynamic if (item[0], item[2][:80]) not in KNOWN_DYNAMIC and (item[0], item[2]) not in KNOWN_DYNAMIC]
    assert not unknown, f"new messages built at run time (list them in KNOWN_DYNAMIC or make them literal): {unknown}"


def test_the_most_specific_pattern_wins():
    assert translate_message("Title can be at most 150 characters.", TAGLISH) == "Ang title ay hanggang 150 na characters lang."
    assert translate_message("Section 2, question 3: type the question.", TAGLISH) == "Section 2, tanong 3: i-type ang tanong."
    assert translate_message("Last name is required.", TAGLISH) == "Kailangan ang apelyido."


def test_word_import_problems_follow_the_language_cookie(teacher):
    from docx import Document
    import io

    document = Document()
    document.add_heading("I. Lesson", level=1)
    document.add_paragraph("Type: Multiple Choice")
    document.add_paragraph("1. A question")
    data = io.BytesIO()
    document.save(data)
    form = {"file": (io.BytesIO(data.getvalue()), "las.docx"), "title": "Philo", "subject": "Philosophy", "dry_run": "1"}
    english = teacher.post("/api/activities/import", data=dict(form, file=(io.BytesIO(data.getvalue()), "las.docx")),
                           content_type="multipart/form-data").get_json()["errors"]
    teacher.set_cookie("myguro-lang", "tl")
    taglish = teacher.post("/api/activities/import", data=dict(form, file=(io.BytesIO(data.getvalue()), "las.docx")),
                           content_type="multipart/form-data").get_json()["errors"]
    assert english and len(english) == len(taglish)
    assert english[0]["message"] != taglish[0]["message"]  # the message is in Taglish, the data around it is the same
    assert english[0]["row"] == taglish[0]["row"]
    assert any("sagot" in problem["message"].lower() or "ANSWER KEY" in problem["message"] for problem in taglish)
