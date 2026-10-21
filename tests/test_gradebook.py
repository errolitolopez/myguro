"""The class gradebook: a row per student, a column per activity, averages and the exports."""

import io

from openpyxl import load_workbook

from conftest import assign, login, make_activity, make_class


def setup_book(teacher):
    class_id, students = make_class(teacher)
    activity_id, q = make_activity(teacher)
    assignment_id = assign(teacher, activity_id, class_id)
    second_id, q2 = make_activity(teacher, title="Second quiz")
    second_assignment = assign(teacher, second_id, class_id)
    return class_id, students, q, assignment_id, second_assignment


def test_gradebook_cells_and_averages(teacher):
    class_id, students, q, assignment_id, second_assignment = setup_book(teacher)
    first = login(students[0]["username"])
    # 1 (MC) + 2 (identification) + 1 of 2 blanks = 4 points; the short answer (5) waits for the teacher
    first.post(f"/api/my-activities/{assignment_id}/submit", json={"answers": {
        str(q[0]): 1, str(q[2]): "Daoism", str(q[3]): ["India", "Japan"], str(q[4]): "Because"}})
    login(students[1]["username"]).put(f"/api/my-activities/{second_assignment}/answers",
                                       json={"answers": {str(q[0]): 1}})

    book = teacher.get(f"/api/classes/{class_id}/gradebook").get_json()
    assert [a["title"] for a in book["activities"]] == ["Early Eastern Philosophy", "Second quiz"]
    assert len(book["students"]) == len(students)
    by_id = {student["id"]: student for student in book["students"]}

    done = by_id[students[0]["id"]]
    assert done["cells"][0]["earned"] == 4 and done["cells"][0]["pending"] == 5
    assert done["cells"][0]["percent"] == 67  # 4 of the 6 points that can be checked so far
    assert done["cells"][1]["status"] == "not_started"
    assert done["submitted"] == 1 and done["average"] == 67

    answering = by_id[students[1]["id"]]
    assert answering["cells"][1]["status"] == "in_progress" and answering["average"] is None
    assert book["activities"][0]["average"] == 67 and book["activities"][1]["average"] is None


def test_excluded_and_not_given_cells(teacher):
    class_id, students = make_class(teacher)
    activity_id, _ = make_activity(teacher)
    chosen = students[0]["id"]
    response = teacher.put(f"/api/activities/{activity_id}/assignments", json={"assignments": [
        {"class_id": class_id, "everyone": False, "student_ids": [chosen, students[1]["id"]]}]})
    assignment_id = next(item["assignment"]["id"] for item in response.get_json()["classes"] if item["id"] == class_id)
    teacher.put(f"/api/assignments/{assignment_id}/students/{students[1]['id']}/excluded", json={"excluded": True})

    book = teacher.get(f"/api/classes/{class_id}/gradebook").get_json()
    by_id = {student["id"]: student for student in book["students"]}
    assert by_id[chosen]["cells"][0]["status"] == "not_started"
    assert by_id[students[1]["id"]]["cells"][0]["status"] == "excluded"
    assert by_id[students[2]["id"]]["cells"][0] is None  # not given this activity


def test_gradebook_is_the_teachers_own_and_exports_work(teacher, other_teacher):
    class_id, students, q, assignment_id, _ = setup_book(teacher)
    login(students[0]["username"]).post(f"/api/my-activities/{assignment_id}/submit",
                                        json={"answers": {str(q[0]): 1}})
    assert other_teacher.get(f"/api/classes/{class_id}/gradebook").status_code == 404
    assert login(students[0]["username"]).get(f"/api/classes/{class_id}/gradebook").status_code == 403

    assert teacher.get(f"/classes/{class_id}/gradebook").status_code == 200
    excel = teacher.get(f"/classes/{class_id}/gradebook.xlsx")
    assert excel.status_code == 200 and "Gradebook - Grade 9 - Rizal - Philosophy.xlsx" in \
        excel.headers["Content-Disposition"]
    sheet = load_workbook(io.BytesIO(excel.data)).active
    rows = [[cell.value for cell in row] for row in sheet.iter_rows()]
    header = next(row for row in rows if row[0] == "No.")
    assert header[5].startswith("Early Eastern Philosophy") and header[-1] == "Average %"
    student_rows = [row for row in rows if isinstance(row[0], int)]
    assert len(student_rows) == len(students)
    assert [row[5] for row in student_rows if row[1] == students[0]["name"]] == [1]  # points as a number

    pdf = teacher.get(f"/classes/{class_id}/gradebook.pdf")
    assert pdf.status_code == 200 and pdf.data.startswith(b"%PDF")
    assert teacher.get(f"/classes/{class_id}/gradebook.csv").status_code == 302
    assert other_teacher.get(f"/classes/{class_id}/gradebook.xlsx").status_code == 302
