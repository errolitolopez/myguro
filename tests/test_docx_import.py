"""Teacher: making an activity sheet from a Word (.docx) file that follows our template."""

import io

from docx import Document

from conftest import error

TITLE = [("p", "LEARNING ACTIVITY SHEET"), ("p", "Name of Student: ____")]


def docx_bytes(lines):
    document = Document()
    for kind, text in lines:
        if kind == "table":
            table = document.add_table(rows=len(text), cols=len(text[0]))
            for row, cells in zip(table.rows, text):
                for cell, value in zip(row.cells, cells):
                    cell.text = value
        elif kind in ("h1", "h2"):
            document.add_heading(text, level=int(kind[1]))
        else:
            document.add_paragraph(text)
    data = io.BytesIO()
    document.save(data)
    return data.getvalue()


def upload(teacher, content, dry_run=True, title="Philo", name="las.docx"):
    return teacher.post("/api/activities/import", data={
        "file": (io.BytesIO(content), name), "title": title, "subject": "Philosophy",
        "dry_run": "1" if dry_run else "0"}, content_type="multipart/form-data")


GOOD = TITLE + [
    ("h1", "I. Lesson"), ("p", "Read this."), ("table", [["Card 1", "Card 2"]]), ("p", "More."),
    ("h1", "II. Pre-test"), ("p", "Type: Multiple Choice"), ("p", "Directions: Choose."),
    ("p", "___ 1. Where did Hinduism start?"), ("p", "A. China"), ("p", "B. India"), ("p", "C. Greece"),
    ("p", "___ 2. Buddhism is about?"), ("p", "A. Suffering"), ("p", "B. War"),
    ("h1", "III. True or False"), ("p", "Type: True or False"), ("p", "Points: 2"),
    ("p", "1. Wu wei is doing nothing."), ("p", "2. Karma is action."),
    ("h1", "IV. Guide"), ("p", "Type: Identification"),
    ("p", "1. Which card is about harmony? ________________"),
    ("h1", "V. Blanks"), ("p", "Type: Fill in the Blank"), ("p", "1. ___ is in ___."),
    ("h1", "VI. Match"), ("p", "Type: Matching"), ("p", "Directions: Match them."),
    ("p", "1. Ren"), ("p", "2. Dharma"), ("p", "A. Duty"), ("p", "B. Humaneness"), ("p", "C. Extra"),
    ("h1", "VII. Reflect"), ("p", "Type: Short Answer"),
    ("p", "1. Why? Answer: __________________"),
    ("h1", "ANSWER KEY"),
    ("p", "II. Pre-test"), ("p", "1. B (India)"), ("p", "2. A"),
    ("h2", "III. True or False"), ("p", "1. FALSE (it means non-forcing)"), ("p", "2. TRUE"),
    ("h2", "IV. Guide"), ("p", "1. Card 3 (Taoism) / Card three"),
    ("h2", "V. Blanks"), ("p", "1. Manila / City of Manila; Philippines (country)"),
    ("h2", "VI. Match"), ("p", "1. B (Humaneness)"), ("p", "2. A (Duty)"),
]


def without(lines, text):
    return [line for line in lines if line[1] != text]


def test_template_download_imports_cleanly(teacher):
    response = teacher.get("/activities/import-template.docx")
    assert response.status_code == 200
    data = upload(teacher, response.data).get_json()
    assert data["errors"] == []
    assert [s["kind"] for s in data["sections"]] == ["text"] + ["activity"] * 6


def test_preview_then_import(teacher):
    data = upload(teacher, docx_bytes(GOOD)).get_json()
    assert data["errors"] == [] and data["imported"] is False
    assert [(s["kind"], s["title"], s["questions"], s["points"]) for s in data["sections"]] == [
        ("text", "I. Lesson", 0, 0), ("activity", "II. Pre-test", 2, 2), ("activity", "III. True or False", 2, 4),
        ("activity", "IV. Guide", 1, 1), ("activity", "V. Blanks", 1, 1), ("activity", "VI. Match", 1, 1),
        ("activity", "VII. Reflect", 1, 5)]
    assert teacher.get("/api/activities").get_json()["activities"] == []

    done = upload(teacher, docx_bytes(GOOD), dry_run=False).get_json()
    assert done["imported"] is True
    sheet = teacher.get(f"/api/activities/{done['activity']['id']}").get_json()["activity"]
    sections = sheet["sections"]
    assert sections[0]["description"] == "Read this.\nCard 1 | Card 2\nMore."
    mc, tf, ident, blank, match, short = (s["questions"] for s in sections[1:])
    assert (mc[0]["prompt"], mc[0]["choices"], mc[0]["answer"]) == ("Where did Hinduism start?", ["China", "India", "Greece"], 1)
    assert [q["answer"] for q in tf] == [False, True]
    assert ident[0]["prompt"] == "Which card is about harmony?"
    assert ident[0]["answer"] == ["Card 3", "Card three"]
    assert blank[0]["answer"] == [["Manila", "City of Manila"], ["Philippines"]]
    assert match[0]["choices"] == ["Duty", "Humaneness", "Extra"]
    assert [(m["item"], m["match"]) for m in match[0]["answer"]] == [("Ren", 1), ("Dharma", 0)]
    assert short[0]["prompt"] == "Why?" and short[0]["points"] == 5


def test_problems_by_section_and_item(teacher):
    lines = without(without(GOOD, "2. A"), "1. FALSE (it means non-forcing)")
    errors = upload(teacher, docx_bytes(lines)).get_json()["errors"]
    assert [(e["where"], e["message"]) for e in errors] == [
        ("II. Pre-test, item 2", "No answer in the ANSWER KEY."),
        ("III. True or False, item 1", "No answer in the ANSWER KEY.")]


def test_missing_type_and_missing_key(teacher):
    lines = [("h1", "II. Quiz"), ("p", "Type: Riddle"), ("p", "1. ?")]
    errors = upload(teacher, docx_bytes(lines)).get_json()["errors"]
    assert errors[0]["where"] == "II. Quiz" and "Type must be one of" in errors[0]["message"]

    no_key = GOOD[:GOOD.index(("h1", "ANSWER KEY"))]
    errors = upload(teacher, docx_bytes(no_key)).get_json()["errors"]
    assert errors and all("no ANSWER KEY" in e["message"] for e in errors)


def test_bad_answers_and_stray_lines(teacher):
    lines = [("h1", "II. Quiz"), ("p", "Type: Multiple Choice"), ("p", "1. Q?"), ("p", "A. x"), ("p", "B. y"),
             ("p", "wandering text"), ("p", "2. Q2?"), ("p", "A. x"), ("p", "B. y"),
             ("h1", "ANSWER KEY"), ("h2", "II. Quiz"), ("p", "1. A"), ("p", "2. Z"), ("p", "9. A")]
    messages = [(e["where"], e["message"]) for e in upload(teacher, docx_bytes(lines)).get_json()["errors"]]
    assert ("II. Quiz, item 1", 'Unexpected text after this item: "wandering text". '
                                "Every question must start with its number, like 3.") in messages
    assert any(where == "II. Quiz, item 2" and "letter of a choice" in text for where, text in messages)
    assert ("II. Quiz, item 9", "The ANSWER KEY has this item but the activity does not.") in messages


def test_not_a_docx_and_wrong_extension(teacher):
    response = upload(teacher, b"not a zip")
    assert response.status_code == 400
    assert "Word .docx" in error(response)
    assert upload(teacher, b"x", name="las.pdf").status_code == 400


def test_admin_cannot_import(admin):
    assert upload(admin, docx_bytes(GOOD)).status_code in (401, 403)
