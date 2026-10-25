"""Teacher: results, scoring, Give another try, paper scores, exports."""

import io

from openpyxl import load_workbook

from conftest import assign, error, login, make_activity, make_class


def setup(teacher, release=True):
    class_id, students = make_class(teacher)
    activity_id, questions = make_activity(teacher)
    assignment_id = assign(teacher, activity_id, class_id, due_date="2020-01-01")
    if release:
        assert teacher.put(f"/api/assignments/{assignment_id}/release", json={"released": True}).status_code == 200
    return assignment_id, activity_id, questions, students


def submit(student_username, assignment_id, answers):
    return login(student_username).post(f"/api/my-activities/{assignment_id}/submit",
                                        json={"answers": {str(key): value for key, value in answers.items()}})


def test_results_table_and_summary(teacher, other_teacher):
    assignment_id, _, q, students = setup(teacher)
    submit(students[0]["username"], assignment_id, {q[0]: 1, q[2]: "Daoism", q[4]: "Because"})
    login(students[1]["username"]).put(f"/api/my-activities/{assignment_id}/answers", json={"answers": {str(q[0]): 0}})
    data = teacher.get(f"/api/assignments/{assignment_id}/results").get_json()
    assert [section["title"] for section in data["sections"]] == ["III. Pre-test", "XI. Reflection"]
    assert data["summary"] == {"students": 6, "submitted": 1, "not_started": 4, "late": 1, "to_check": 5,
                               "average": 50, "excluded": 0}
    rows = {row["id"]: row for row in data["students"]}
    assert rows[students[0]["id"]]["status"] == "submitted" and rows[students[1]["id"]]["status"] == "in_progress"
    assert other_teacher.get(f"/api/assignments/{assignment_id}/results").status_code == 404


def test_score_short_answer_and_override(teacher, other_teacher):
    assignment_id, _, q, students = setup(teacher)
    submit(students[0]["username"], assignment_id, {q[0]: 1, q[2]: "Dao", q[4]: "Because"})
    detail = teacher.get(f"/api/assignments/{assignment_id}/students/{students[0]['id']}").get_json()["detail"]
    points = lambda values: teacher.put(f"/api/submissions/{detail['id']}/points",
                                        json={"points": {str(key): value for key, value in values.items()}})
    assert points({q[4]: 4, q[2]: 2}).status_code == 200
    assert error(points({q[4]: 6})) == "Points must be a whole number from 0 to 5."
    assert error(points({q[0]: None})) == "Points must be a whole number from 0 to 1."
    assert other_teacher.put(f"/api/submissions/{detail['id']}/points", json={"points": {}}).status_code == 404
    student = login(students[0]["username"])
    assert student.get(f"/api/my-activities/{assignment_id}").get_json()["activity"]["score"]["total"] == \
        {"earned": 7, "pending": 0, "possible": 11}


def test_give_another_try_and_short_answer_carry_over(teacher):
    assignment_id, _, q, students = setup(teacher)
    username, student_id = students[0]["username"], students[0]["id"]
    submit(username, assignment_id, {q[0]: 1, q[4]: "Same text"})
    detail = teacher.get(f"/api/assignments/{assignment_id}/students/{student_id}").get_json()["detail"]
    teacher.put(f"/api/submissions/{detail['id']}/points", json={"points": {str(q[4]): 3}})
    reopened = teacher.post(f"/api/assignments/{assignment_id}/students/{student_id}/reopen").get_json()
    assert [attempt["attempt"] for attempt in reopened["attempts"]] == [1, 2]
    assert error(teacher.post(f"/api/assignments/{assignment_id}/students/{student_id}/reopen")) == \
        "The student is already answering (not submitted yet)."
    activity = submit(username, assignment_id, {}).get_json()["activity"]
    assert activity["score"]["total"] == {"earned": 4, "pending": 0, "possible": 11}  # 1 + the kept 3


def test_paper_score(teacher):
    assignment_id, activity_id, q, students = setup(teacher)
    teacher.put(f"/api/activities/{activity_id}/settings", json={"max_attempts": None})
    student_id, username = students[0]["id"], students[0]["username"]
    url = f"/api/assignments/{assignment_id}/students/{student_id}/paper"
    assert error(teacher.post(url, json={"score": 12})) == "The score must be a whole number from 0 to 11."
    data = teacher.post(url, json={"score": 9, "late": True, "note": "Passed Oct 14"}).get_json()
    assert data["detail"]["paper"] is True and data["detail"]["score"]["total"]["earned"] == 9
    student = login(username)
    activity = student.get(f"/api/my-activities/{assignment_id}").get_json()["activity"]
    assert activity["paper"] is True and activity["can_retry"] is False and activity["note"] == "Passed Oct 14"
    assert error(student.post(f"/api/my-activities/{assignment_id}/retry")) == \
        "Your teacher already scored this from your paper."
    assert error(student.put(f"/api/my-activities/{assignment_id}/answers", json={"answers": {}})) == \
        "You already submitted this activity."
    paper_id = data["detail"]["id"]
    assert teacher.put(f"/api/submissions/{paper_id}/paper", json={"score": 10}).get_json()["detail"]["score"][
        "total"]["earned"] == 10
    assert error(teacher.put(f"/api/submissions/{paper_id}/points", json={"points": {}})) == \
        "This is a paper score: change its total instead."
    assert teacher.delete(f"/api/submissions/{paper_id}/paper").get_json()["attempts"] == []


def test_exports(teacher, other_teacher):
    assignment_id, _, q, students = setup(teacher)
    submit(students[0]["username"], assignment_id, {q[0]: 1})
    xlsx = teacher.get(f"/results/{assignment_id}/export.xlsx")
    assert xlsx.status_code == 200 and "Early Eastern Philosophy - Grade 9 - Rizal - Philosophy.xlsx" in \
        xlsx.headers["Content-Disposition"]
    rows = list(load_workbook(io.BytesIO(xlsx.data)).active.iter_rows(values_only=True))
    assert rows[0][0] == "Early Eastern Philosophy" and rows[5][:3] == ("No.", "Student", "Username")
    assert len(rows) == 6 + 6 + 2  # the table, a blank row and the "Generated by" note
    assert rows[-1][0] == "Generated by MyGuro \u2013 Learning Management System"
    pdf = teacher.get(f"/results/{assignment_id}/export.pdf")
    assert pdf.status_code == 200 and pdf.data.startswith(b"%PDF")
    assert other_teacher.get(f"/results/{assignment_id}/export.pdf").status_code == 302
    assert teacher.get(f"/results/{assignment_id}/export.doc").status_code == 302


def test_check_card_lists_every_class_the_sheet_is_given_to(teacher, other_teacher):
    assignment_id, activity_id, q, students = setup(teacher)
    card = teacher.get(f"/api/results/to-check?activity_id={activity_id}").get_json()["assignments"][0]
    assert card["id"] == assignment_id and card["submitted"] == 0 and card["average"] is None  # nobody yet
    submit(students[0]["username"], assignment_id, {q[0]: 1, q[2]: "Daoism", q[4]: "Because"})
    card = teacher.get(f"/api/results/to-check?activity_id={activity_id}").get_json()["assignments"][0]
    assert card["submitted"] == 1 and card["to_check"] == 5 and card["due_date"] == "2020-01-01"
    assert teacher.get("/api/results/to-check?activity_id=99999").get_json() == {"assignments": []}
    assert other_teacher.get("/api/results/to-check").get_json() == {"assignments": []}


def test_results_and_the_checking_page_tell_when_the_student_started_and_submitted(teacher):
    class_id, students = make_class(teacher)
    activity_id, q = make_activity(teacher)
    assignment_id = assign(teacher, activity_id, class_id)
    student = login(students[0]["username"])
    student.put(f"/api/my-activities/{assignment_id}/answers", json={"answers": {str(q[0]): 1}})

    row = next(r for r in teacher.get(f"/api/assignments/{assignment_id}/results").get_json()["students"]
               if r["id"] == students[0]["id"])
    assert row["started_at"] and row["submitted_at"] is None  # still answering
    other = next(r for r in teacher.get(f"/api/assignments/{assignment_id}/results").get_json()["students"]
                 if r["id"] == students[1]["id"])
    assert other["started_at"] is None

    student.post(f"/api/my-activities/{assignment_id}/submit", json={"answers": {str(q[0]): 1}})
    detail = teacher.get(f"/api/assignments/{assignment_id}/students/{students[0]['id']}").get_json()["detail"]
    assert detail["started_at"] and detail["submitted_at"] and detail["auto_submitted"] is False
