"""Teacher: making an activity sheet from an Excel file."""

import io

from openpyxl import Workbook, load_workbook

from conftest import error

COLUMNS = ("Row", "Title", "Description", "Question type",
           "Choice A", "Choice B", "Choice C", "Choice D", "Choice E", "Choice F", "Answer", "Points")


def xlsx(rows, headers=COLUMNS):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(list(headers))
    for row in rows:
        sheet.append(list(row))
    data = io.BytesIO()
    workbook.save(data)
    return data.getvalue()


def question(kind, prompt, answer, choices=(), points=None):
    padded = (list(choices) + [None] * 6)[:6]
    return ("question", None, prompt, kind, *padded, answer, points)


def upload(teacher, content, dry_run, title="Cells", subject="Science"):
    return teacher.post("/api/activities/import", data={
        "file": (io.BytesIO(content), "cells.xlsx"), "title": title, "subject": subject,
        "dry_run": "1" if dry_run else "0"}, content_type="multipart/form-data")


GOOD = [
    ("text", "I. Lesson", "Read this.", None, None, None, None, None, None, None, None, None),
    ("activity", "II. Check", "Answer all.", None, None, None, None, None, None, None, None, None),
    question("multiple_choice", "2 + 2?", "b", ["3", "4", "5"]),
    question("true_false", "Sky is blue", "True", points=2),
    question("identification", "Capital?", "Manila / City of Manila", points=3),
    question("fill_blank", "___ is the capital of ___", "Manila ; Philippines / PH"),
    question("short_answer", "Explain.", None),
]


def test_template_download(teacher):
    response = teacher.get("/activities/import-template.xlsx")
    assert response.status_code == 200
    sheet = load_workbook(io.BytesIO(response.data)).worksheets[0]
    assert tuple(cell.value for cell in sheet[1]) == COLUMNS


def test_preview_then_import(teacher):
    preview = upload(teacher, xlsx(GOOD), dry_run=True)
    assert preview.status_code == 200
    data = preview.get_json()
    assert data["errors"] == [] and data["imported"] is False
    assert [(s["kind"], s["questions"], s["points"]) for s in data["sections"]] == [("text", 0, 0), ("activity", 5, 12)]
    assert teacher.get("/api/activities").get_json()["activities"] == []  # nothing created yet

    done = upload(teacher, xlsx(GOOD), dry_run=False).get_json()
    assert done["imported"] is True
    sheet = teacher.get(f"/api/activities/{done['activity']['id']}").get_json()["activity"]
    assert (sheet["title"], sheet["subject"], sheet["assigned"]) == ("Cells", "Science", 0)
    assert [s["kind"] for s in sheet["sections"]] == ["text", "activity"]
    questions = sheet["sections"][1]["questions"]
    assert [(q["type"], q["points"]) for q in questions] == [
        ("multiple_choice", 1), ("true_false", 2), ("identification", 3), ("fill_blank", 1), ("short_answer", 5)]
    assert questions[0]["choices"] == ["3", "4", "5"] and questions[0]["answer"] == 1
    assert questions[1]["answer"] is True
    assert questions[2]["answer"] == ["Manila", "City of Manila"]
    assert questions[3]["answer"] == [["Manila"], ["Philippines", "PH"]]


def test_problems_listed_with_row_numbers_and_nothing_created(teacher):
    content = xlsx([
        ("heading", "x", "y", None, None, None, None, None, None, None, None, None),           # row 2
        question("true_false", "Orphan", "True"),                                              # row 3
        ("activity", "Part", "Directions", None, None, None, None, None, None, None, None, None),
        question("multiple_choice", "Pick", "C", ["a", "b"]),                                  # row 5
        question("riddle", "?", "x"),                                                          # row 6
        question("fill_blank", "No blank here", "x"),                                          # row 7
        question("identification", "Capital?", "Manila", points="500"),                        # row 8
        ("activity", "Empty part", "Nothing", None, None, None, None, None, None, None, None, None),  # row 9
    ])
    data = upload(teacher, content, dry_run=False).get_json()
    assert data["imported"] is False
    assert sorted(e["row"] for e in data["errors"]) == [2, 3, 5, 6, 7, 8]
    # Row 9 is only reported once the row-level problems are fixed
    assert teacher.get("/api/activities").get_json()["activities"] == []


def test_activity_without_questions(teacher):
    content = xlsx([("activity", "Part", "Directions", None, None, None, None, None, None, None, None, None)])
    data = upload(teacher, content, dry_run=True).get_json()
    assert [e["row"] for e in data["errors"]] == [2]


def test_bad_files_and_title(teacher):
    assert upload(teacher, b"not excel", dry_run=True).status_code == 400
    assert upload(teacher, xlsx([], headers=("Foo", "Bar")), dry_run=True).status_code == 400
    assert "no rows" in error(upload(teacher, xlsx([]), dry_run=True))
    assert upload(teacher, xlsx(GOOD), dry_run=True, title="  ").status_code == 400


def test_teachers_only(admin):
    assert upload(admin, xlsx(GOOD), dry_run=True).status_code in (401, 403)


def test_matching_rows(teacher):
    rows = [
        ("activity", "Match", "Directions.", None, None, None, None, None, None, None, None, None),
        question("matching", "Match the country with its capital.",
                 "Philippines = Manila ; Japan = Tokyo ; Thailand = bangkok", ["Seoul", "tokyo"], points=6),
    ]
    data = upload(teacher, xlsx(rows), dry_run=False).get_json()
    assert data["errors"] == [] and data["imported"] is True
    activity = teacher.get(f"/api/activities/{data['activity']['id']}").get_json()["activity"]
    matching = activity["sections"][0]["questions"][0]
    assert matching["type"] == "matching" and matching["points"] == 6
    assert matching["choices"] == ["bangkok", "Manila", "Seoul", "Tokyo"]  # sorted, Seoul is an extra
    assert [(pair["item"], matching["choices"][pair["match"]]) for pair in matching["answer"]] == \
        [("Philippines", "Manila"), ("Japan", "Tokyo"), ("Thailand", "bangkok")]


def test_matching_problems(teacher):
    rows = [
        ("activity", "A", "Directions.", None, None, None, None, None, None, None, None, None),
        question("matching", "No pairs", "Manila"),
        question("matching", "One pair", "Japan = Tokyo"),
        question("matching", "Empty side", "Japan = ; Peru = Lima"),
    ]
    errors = upload(teacher, xlsx(rows), dry_run=True).get_json()["errors"]
    pairs = 'Answer must be pairs like "Hinduism = India ; Taoism = China".'
    assert [(e["row"], e["message"]) for e in errors] == [
        (3, pairs), (4, "Matching needs at least 2 pairs, split by ;."), (5, pairs)]
