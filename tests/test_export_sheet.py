"""Teacher: download a sheet with its answers as Word or Excel, and import the Word file again as a new activity."""

import io

from openpyxl import load_workbook

from conftest import make_activity

SECTIONS = [
    {"kind": "text", "title": "I. Lesson", "description": "<p>Plants <b>make</b> food.</p><ul><li>Light</li></ul>",
     "description_format": "rich"},
    {"kind": "activity", "title": "II. Quiz", "description": "Answer all.", "questions": [
        {"type": "multiple_choice", "prompt": "Red planet?", "choices": ["Venus", "Mars", "Earth"],
         "answer": 1, "points": 2},
        {"type": "multiple_choice", "prompt": "Biggest planet?", "choices": ["Jupiter", "Mars"],
         "answer": 0, "points": 2},
        {"type": "true_false", "prompt": "Water boils at 100 C.", "answer": True, "points": 1},
        {"type": "identification", "prompt": "Gas plants take in?", "answer": ["carbon dioxide", "CO2"], "points": 1},
        {"type": "fill_blank", "prompt": "The ___ orbits the ___.", "answer": [["moon"], ["earth", "Earth"]],
         "points": 2},
        {"type": "matching", "prompt": "Match them.", "choices": ["Star", "Planet", "Moon", "Comet"],
         "answer": [{"item": "Sun", "match": 0}, {"item": "Mars", "match": 1}], "points": 3},
        {"type": "short_answer", "prompt": "Explain photosynthesis.", "answer": "light and CO2", "points": 5},
    ]},
    {"kind": "activity", "title": "II. Quiz", "questions": [  # the same title again
        {"type": "true_false", "prompt": "The sky is green.", "answer": False, "points": 1},
    ]},
]


def comparable(activity):
    """What must survive a round trip: every question, in order (matching by its pairs, since the
    import sorts Column B)."""
    found = []
    for section in activity["sections"]:
        for question in section["questions"]:
            answer, choices = question["answer"], question["choices"]
            if question["type"] == "matching":
                answer = sorted((pair["item"], choices[pair["match"]]) for pair in answer)
                choices = sorted(choices)
            found.append((question["type"], question["prompt"], choices, answer))
    return found


def round_trip(teacher, activity_id, kind):
    download = teacher.get(f"/activities/{activity_id}/export.{kind}")
    assert download.status_code == 200
    assert download.headers["Content-Disposition"].startswith("attachment")
    response = teacher.post("/api/activities/import", data={
        "file": (io.BytesIO(download.data), f"sheet.{kind}"), "title": "Again", "dry_run": "0"},
        content_type="multipart/form-data")
    assert response.status_code == 200 and response.get_json()["imported"], response.get_json()
    new_id = response.get_json()["activity"]["id"]
    return teacher.get(f"/api/activities/{new_id}").get_json()["activity"], new_id


def test_download_then_import_gives_the_same_questions(teacher):
    activity_id, _ = make_activity(teacher, sections=SECTIONS)
    original = teacher.get(f"/api/activities/{activity_id}").get_json()["activity"]
    again, new_id = round_trip(teacher, activity_id, "docx")
    assert new_id != activity_id
    assert comparable(again) == comparable(original)
    assert again["sections"][0]["title"] == "I. Lesson"
    assert "Plants make food." in again["sections"][0]["description"]


def test_excel_keeps_points_per_question(teacher):
    activity_id, _ = make_activity(teacher, sections=SECTIONS)
    download = teacher.get(f"/activities/{activity_id}/export.xlsx")
    assert download.status_code == 200
    rows = list(load_workbook(io.BytesIO(download.data)).worksheets[0].iter_rows(min_row=2, values_only=True))
    assert [row[11] for row in rows if row[0] == "question"] == [2, 2, 1, 1, 2, 3, 5, 1]


def test_word_splits_a_mixed_section_and_keeps_headings_unique(teacher):
    activity_id, _ = make_activity(teacher, sections=SECTIONS)
    again, _ = round_trip(teacher, activity_id, "docx")
    titles = [section["title"] for section in again["sections"]]
    assert len(titles) == len({title.casefold() for title in titles})  # the key can tell them apart
    assert titles[:3] == ["I. Lesson", "II. Quiz", "II. Quiz (part 2)"]


def test_export_is_only_for_the_owner_and_known_formats(teacher, other_teacher):
    activity_id, _ = make_activity(teacher)
    assert other_teacher.get(f"/activities/{activity_id}/export.docx").status_code == 302
    assert teacher.get(f"/activities/{activity_id}/export.pdf").status_code == 302
